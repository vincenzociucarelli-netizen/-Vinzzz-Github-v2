"""Tests for new endpoints:
- GET /api/admin/compensation-summary/default-range
- GET /api/admin/compensation-summary/export-xlsx
"""
import os
import io
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("EXPO_PUBLIC_BACKEND_URL", "https://zaino-coach-hub.preview.emergentagent.com").rstrip("/")

ADMIN_EMAIL = "vincenzo.ciucarelli@gmail.com"
ADMIN_PASSWORD = "Test1234"
NON_ADMIN_EMAIL = "viola@test.com"
NON_ADMIN_PASSWORD = "Test1234"


def _login(email: str, password: str) -> str:
    r = requests.post(
        f"{BASE_URL}/api/auth/login",
        json={"email": email, "password": password},
        timeout=20,
    )
    assert r.status_code == 200, f"login {email} failed: {r.status_code} {r.text[:200]}"
    j = r.json()
    return j.get("access_token") or j.get("token")


@pytest.fixture(scope="module")
def admin_token():
    return _login(ADMIN_EMAIL, ADMIN_PASSWORD)


@pytest.fixture(scope="module")
def non_admin_token():
    return _login(NON_ADMIN_EMAIL, NON_ADMIN_PASSWORD)


# ---------------- default-range ----------------
class TestDefaultRange:
    def test_admin_default_range_returns_season_2025_2026(self, admin_token):
        r = requests.get(
            f"{BASE_URL}/api/admin/compensation-summary/default-range",
            headers={"Authorization": f"Bearer {admin_token}"},
            timeout=15,
        )
        assert r.status_code == 200, r.text
        data = r.json()
        assert data.get("start_date") == "2025-07-01", data
        assert data.get("end_date") == "2026-06-30", data
        assert data.get("label") == "Stagione 2025/2026", data

    def test_non_admin_forbidden(self, non_admin_token):
        r = requests.get(
            f"{BASE_URL}/api/admin/compensation-summary/default-range",
            headers={"Authorization": f"Bearer {non_admin_token}"},
            timeout=15,
        )
        assert r.status_code == 403, f"expected 403 got {r.status_code} {r.text[:200]}"

    def test_no_auth_returns_401(self):
        r = requests.get(
            f"{BASE_URL}/api/admin/compensation-summary/default-range", timeout=15
        )
        assert r.status_code in (401, 403)


# ---------------- export-xlsx ----------------
class TestExportXLSX:
    def test_export_xlsx_success(self, admin_token):
        r = requests.get(
            f"{BASE_URL}/api/admin/compensation-summary/export-xlsx",
            params={"start_date": "2025-09-01", "end_date": "2026-06-30"},
            headers={"Authorization": f"Bearer {admin_token}"},
            timeout=60,
        )
        assert r.status_code == 200, r.text[:500]
        # content-type
        ctype = r.headers.get("content-type", "")
        assert "spreadsheetml.sheet" in ctype, f"unexpected content-type: {ctype}"
        # parse XLSX
        wb = load_workbook(io.BytesIO(r.content), data_only=False)
        ws = wb.active
        # Find header row (the row that contains 'Cognome')
        header_row_idx = None
        for row in ws.iter_rows(min_row=1, max_row=10, values_only=False):
            for cell in row:
                if cell.value and "Cognome" in str(cell.value):
                    header_row_idx = cell.row
                    break
            if header_row_idx:
                break
        assert header_row_idx is not None, "Header row 'Cognome' not found"
        headers = [c.value for c in ws[header_row_idx]]
        print("HEADERS:", headers)
        # Required columns
        assert any(h and "Totale" in str(h) and "Range" in str(h) for h in headers), headers
        assert any(h and "Esterni" in str(h) for h in headers), headers
        # Data rows: collect all non-empty rows after header until TOTALI
        data_rows = []
        totali_row = None
        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=False):
            first = row[0].value
            if first and "TOTALI" in str(first).upper():
                totali_row = row
                break
            if any(c.value not in (None, "") for c in row):
                data_rows.append(row)
        print(f"data rows count = {len(data_rows)}")
        assert len(data_rows) >= 10, f"expected ~10+ people, found {len(data_rows)}"

        # find col indexes
        def col_idx(predicate):
            for i, h in enumerate(headers, start=1):
                if h and predicate(str(h)):
                    return i
            return None

        col_total = col_idx(lambda h: "Totale" in h and "Range" in h)
        col_external = col_idx(lambda h: "Esterni" in h)
        col_lastname = col_idx(lambda h: "Cognome" in h)
        col_firstname = col_idx(lambda h: "Nome" in h and "Cognome" not in h)
        col_5k = col_idx(lambda h: "5.000" in h or "5K" in h or "5000" in h or "€5" in h)
        col_15k = col_idx(lambda h: "15.000" in h or "15K" in h or "15000" in h or "€15" in h)
        assert col_total, "Totale range column not found"
        assert col_external, "Esterni column not found"

        # Look for Fede Rica with expected 17000
        fede_total = None
        for row in data_rows:
            ln = row[col_lastname - 1].value if col_lastname else None
            fn = row[col_firstname - 1].value if col_firstname else None
            if ln and fn and "rica" in str(ln).lower() and "fede" in str(fn).lower():
                fede_total = row[col_total - 1].value
                break
        if fede_total is not None:
            assert abs(float(fede_total) - 17000.0) < 1.0, f"Fede Rica total expected 17000, got {fede_total}"
            print(f"Fede Rica total OK: {fede_total}")

        # TOTALI row should contain SUM formula
        if totali_row is not None:
            sum_found = False
            for c in totali_row:
                if c.value and isinstance(c.value, str) and c.value.upper().startswith("=SUM"):
                    sum_found = True
                    break
            assert sum_found, "TOTALI row must contain a =SUM formula"

        # threshold flag column present
        assert col_5k or col_15k, "expected at least one threshold column (5K or 15K)"

    def test_export_invalid_date_range(self, admin_token):
        r = requests.get(
            f"{BASE_URL}/api/admin/compensation-summary/export-xlsx",
            params={"start_date": "2026-06-30", "end_date": "2025-07-01"},
            headers={"Authorization": f"Bearer {admin_token}"},
            timeout=20,
        )
        assert r.status_code == 400, f"expected 400 got {r.status_code} {r.text[:200]}"

    def test_export_invalid_date_format(self, admin_token):
        r = requests.get(
            f"{BASE_URL}/api/admin/compensation-summary/export-xlsx",
            params={"start_date": "not-a-date", "end_date": "2026-06-30"},
            headers={"Authorization": f"Bearer {admin_token}"},
            timeout=20,
        )
        assert r.status_code == 400, f"expected 400 got {r.status_code}"

    def test_export_non_admin_forbidden(self, non_admin_token):
        r = requests.get(
            f"{BASE_URL}/api/admin/compensation-summary/export-xlsx",
            params={"start_date": "2025-07-01", "end_date": "2026-06-30"},
            headers={"Authorization": f"Bearer {non_admin_token}"},
            timeout=20,
        )
        assert r.status_code == 403
